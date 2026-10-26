import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(target_path="Golden_Snitch_Pitch_Deck_Enhanced.pptx"):
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # APPROVED COLOR SYSTEM — RESTRAINED PALETTE
    C_BG = RGBColor(16, 20, 29)            # Deep charcoal #10141D
    C_SURFACE = RGBColor(25, 31, 43)        # Elevated charcoal #191F2B
    C_SURFACE_LIGHT = RGBColor(34, 42, 56)  # Secondary surface #222A38
    C_TEXT_PRI = RGBColor(245, 247, 250)    # Primary white #F5F7FA
    C_TEXT_SEC = RGBColor(180, 190, 204)    # Secondary slate #B4BECC
    C_TEXT_MUTED = RGBColor(132, 145, 163)  # Muted labels #8491A3
    C_GOLD = RGBColor(245, 184, 46)         # Signature gold #F5B82E
    C_CYAN = RGBColor(53, 213, 232)         # Restrained cyan #35D5E8
    C_BORDER = RGBColor(48, 57, 74)         # Subtle border #30394A
    C_RED = RGBColor(241, 107, 114)         # Negative indicator #F16B72

    FONT_FAMILY = "Segoe UI"

    def apply_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title, subtitle):
        tb_t = slide.shapes.add_textbox(Inches(0.70), Inches(0.48), Inches(11.933), Inches(0.48))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_PRI

        tb_s = slide.shapes.add_textbox(Inches(0.70), Inches(1.04), Inches(11.933), Inches(0.35))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle
        p_s.font.name = FONT_FAMILY
        p_s.font.size = Pt(13)
        p_s.font.color.rgb = C_TEXT_SEC

    def add_footer(slide, current_slide, total_slides=6):
        tb_f = slide.shapes.add_textbox(Inches(0.70), Inches(6.92), Inches(11.933), Inches(0.28))
        tf_f = tb_f.text_frame
        tf_f.margin_left = tf_f.margin_right = tf_f.margin_top = tf_f.margin_bottom = 0
        p = tf_f.paragraphs[0]
        p.text = f"Golden Snitch (ScoutIQ)  |  Team Logic Legend  |  Live Platform: https://goldensnitch.up.railway.app  |  Slide {current_slide} of {total_slides}"
        p.font.name = FONT_FAMILY
        p.font.size = Pt(9.5)
        p.font.color.rgb = C_TEXT_MUTED

    def add_card(slide, left, top, width, height, bg=C_SURFACE, border=C_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        if border:
            card.line.color.rgb = border
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    def set_notes(slide, opening, main_point, evidence, transition):
        tf = slide.notes_slide.notes_text_frame
        tf.text = (
            f"1. OPENING: {opening}\n\n"
            f"2. MAIN POINT: {main_point}\n\n"
            f"3. EVIDENCE: {evidence}\n\n"
            f"4. TRANSITION: {transition}"
        )

    # =========================================================================
    # SLIDE 1 — VISION: RESTYLED WITH GENEROUS SPACING & LARGE HERO UI
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_bg(slide1)

    # Left: Typography & Value Proposition
    tb_b = slide1.shapes.add_textbox(Inches(0.70), Inches(0.80), Inches(5.40), Inches(2.30))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0

    p_tag = tf_b.paragraphs[0]
    p_tag.text = "ESPORTS OPERATIONS INTELLIGENCE"
    p_tag.font.name = FONT_FAMILY
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = C_GOLD

    p_main = tf_b.add_paragraph()
    p_main.text = "Golden Snitch"
    p_main.font.name = FONT_FAMILY
    p_main.font.size = Pt(40)
    p_main.font.bold = True
    p_main.font.color.rgb = C_TEXT_PRI
    p_main.space_before = Pt(4)

    p_sub = tf_b.add_paragraph()
    p_sub.text = "AI-powered esports scouting and multi-roster intelligence platform."
    p_sub.font.name = FONT_FAMILY
    p_sub.font.size = Pt(15)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_TEXT_SEC
    p_sub.space_before = Pt(6)

    p_val = tf_b.add_paragraph()
    p_val.text = "Transforming competitive recruitment from superficial frag hunting into an objective, gap-driven decision science."
    p_val.font.name = FONT_FAMILY
    p_val.font.size = Pt(12)
    p_val.font.color.rgb = C_TEXT_MUTED
    p_val.space_before = Pt(6)

    # 3 Crisp Verification Badges (Generous Spacing)
    badges = [
        ("1,726 VCT MATCHES", "Calibrated on 10 VCT LAN stages with 156 verified pros."),
        ("88% RECRUITMENT SPEEDUP", "Scouting cycle reduced from 18.5 hours to 2.2 hours."),
        ("PRODUCTION READY", "112 passing pytest unit tests, 0 lint errors, live Railway deployment.")
    ]
    tb_badge = slide1.shapes.add_textbox(Inches(0.70), Inches(3.45), Inches(5.40), Inches(3.20))
    tf_badge = tb_badge.text_frame
    tf_badge.word_wrap = True
    tf_badge.margin_left = tf_badge.margin_right = tf_badge.margin_top = tf_badge.margin_bottom = 0

    for idx, (b_title, b_desc) in enumerate(badges):
        p_bh = tf_badge.paragraphs[0] if idx == 0 else tf_badge.add_paragraph()
        p_bh.text = b_title
        p_bh.font.name = FONT_FAMILY
        p_bh.font.size = Pt(12)
        p_bh.font.bold = True
        p_bh.font.color.rgb = C_GOLD if idx == 0 else C_TEXT_PRI
        if idx > 0:
            p_bh.space_before = Pt(14)

        p_bd = tf_badge.add_paragraph()
        p_bd.text = b_desc
        p_bd.font.name = FONT_FAMILY
        p_bd.font.size = Pt(10.5)
        p_bd.font.color.rgb = C_TEXT_MUTED
        p_bd.space_before = Pt(2)

    # Right: Enlarged Product Interface Showcase
    hero_w = Inches(6.30)
    hero_h = Inches(4.35)
    hero_left = Inches(6.33)
    hero_top = Inches(1.30)
    add_card(slide1, hero_left, hero_top, hero_w, hero_h)

    hero_img = "SiteSnipts/02_intelligence_hub.png"
    if os.path.exists(hero_img):
        slide1.shapes.add_picture(hero_img, hero_left + Inches(0.08), hero_top + Inches(0.08), hero_w - Inches(0.16), hero_h - Inches(0.16))

    tb_cap = slide1.shapes.add_textbox(hero_left, hero_top + hero_h + Inches(0.15), hero_w, Inches(0.50))
    tf_cap = tb_cap.text_frame
    tf_cap.word_wrap = True
    tf_cap.margin_left = tf_cap.margin_right = tf_cap.margin_top = tf_cap.margin_bottom = 0
    p_c = tf_cap.paragraphs[0]
    p_c.text = "ScoutIQ Command Center // Live Vacancy Gap Matching"
    p_c.font.name = FONT_FAMILY
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = C_TEXT_PRI

    p_cd = tf_cap.add_paragraph()
    p_cd.text = "Deterministic candidate fit scoring benchmarked under $1.5M roster payroll caps."
    p_cd.font.name = FONT_FAMILY
    p_cd.font.size = Pt(10)
    p_cd.font.color.rgb = C_TEXT_MUTED

    add_footer(slide1, 1)

    set_notes(
        slide1,
        "Good morning members of the jury. We present Golden Snitch: a roster-first esports operations intelligence platform.",
        "Traditional esports scouting burns millions every season by signing high-frag players who don't fit the team's tactical needs.",
        "Golden Snitch is an operational production-ready solution evaluated on 1,726 VCT LAN matches and backed by 112 passing automated test cases.",
        "Let us examine why traditional esports scouting breaks down."
    )

    # =========================================================================
    # SLIDE 2 — PRODUCT-FIRST EVIDENCE: ENLARGED COMPARISON & CRISP METRICS
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_bg(slide2)

    add_header(
        slide2,
        "Traditional scouting misses the players who win rounds.",
        "Raw combat statistics can hide tactical impact, role fit, and decision-making quality."
    )

    # Left: 3 Crisp, Spaced Problem Points (No bulky nested panels!)
    left_w = Inches(4.90)
    tb_pr = slide2.shapes.add_textbox(Inches(0.70), Inches(1.65), left_w, Inches(4.80))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_right = tf_pr.margin_top = tf_pr.margin_bottom = 0

    probs = [
        ("Combat-Stat Bias",
         "Public leaderboards rank players solely on raw combat score (ACS) and K/D ratios. This systematically undervalues support anchors, controllers, and vocal shotcallers."),
        ("Context Blindness",
         "High kill counts often hide pad-stat behavior. In LAN matches, up to 42% of a frag-hunter's frags occur in already-lost save rounds (exit frags), carrying zero round-win equity."),
        ("Recruitment & Buyout Risk",
         "Organizations spend ₹50 Lakhs to ₹1.5 Crores signing 'star' fraggers into support vacancies, leading to severe role collisions, benchings, and financial write-offs.")
    ]
    for idx, (title, body) in enumerate(probs):
        p_t = tf_pr.paragraphs[0] if idx == 0 else tf_pr.add_paragraph()
        p_t.text = f"{idx+1}.  {title}"
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_PRI
        if idx > 0:
            p_t.space_before = Pt(16)

        p_b = tf_pr.add_paragraph()
        p_b.text = body
        p_b.font.name = FONT_FAMILY
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = C_TEXT_SEC
        p_b.space_before = Pt(4)

    # Right: ENLARGED PRODUCT COMPARISON EVIDENCE
    right_w = Inches(6.70)
    right_left = Inches(5.93)
    right_top = Inches(1.65)
    img_h = Inches(2.90)

    # Enlarged Player Dossier Interface
    add_card(slide2, right_left, right_top, right_w, img_h)
    dossier_img = "SiteSnipts/05_player_dossier.png"
    if os.path.exists(dossier_img):
        slide2.shapes.add_picture(dossier_img, right_left + Inches(0.08), right_top + Inches(0.08), right_w - Inches(0.16), img_h - Inches(0.16))

    # Two Aligned Metric High-Contrast Chips
    card_y = right_top + img_h + Inches(0.18)
    card_w = Inches(3.25)
    card_h = Inches(1.80)

    # Voltz (Decoy)
    add_card(slide2, right_left, card_y, card_w, card_h)
    tb_cA = slide2.shapes.add_textbox(right_left + Inches(0.15), card_y + Inches(0.12), card_w - Inches(0.30), card_h - Inches(0.24))
    tf_cA = tb_cA.text_frame
    tf_cA.word_wrap = True
    tf_cA.margin_left = tf_cA.margin_right = tf_cA.margin_top = tf_cA.margin_bottom = 0

    p_a1 = tf_cA.paragraphs[0]
    p_a1.text = "VOLTZ // STATISTICAL DECOY"
    p_a1.font.name = FONT_FAMILY
    p_a1.font.size = Pt(9.5)
    p_a1.font.bold = True
    p_a1.font.color.rgb = C_RED

    p_a2 = tf_cA.add_paragraph()
    p_a2.text = "49.3"
    p_a2.font.name = FONT_FAMILY
    p_a2.font.size = Pt(32)
    p_a2.font.bold = True
    p_a2.font.color.rgb = C_RED
    p_a2.space_before = Pt(1)

    p_a3 = tf_cA.add_paragraph()
    p_a3.text = "285 ACS (#1 in frags), but 42% non-impact exit frags. Fails to cover Sentinels' utility deficit."
    p_a3.font.name = FONT_FAMILY
    p_a3.font.size = Pt(9.5)
    p_a3.font.color.rgb = C_TEXT_MUTED
    p_a3.space_before = Pt(2)

    # Mirage (True Fit)
    add_card(slide2, right_left + card_w + Inches(0.20), card_y, card_w, card_h)
    tb_cB = slide2.shapes.add_textbox(right_left + card_w + Inches(0.35), card_y + Inches(0.12), card_w - Inches(0.30), card_h - Inches(0.24))
    tf_cB = tb_cB.text_frame
    tf_cB.word_wrap = True
    tf_cB.margin_left = tf_cB.margin_right = tf_cB.margin_top = tf_cB.margin_bottom = 0

    p_b1 = tf_cB.paragraphs[0]
    p_b1.text = "MIRAGE // OPTIMAL ROSTER FIT"
    p_b1.font.name = FONT_FAMILY
    p_b1.font.size = Pt(9.5)
    p_b1.font.bold = True
    p_b1.font.color.rgb = C_GOLD

    p_b2 = tf_cB.add_paragraph()
    p_b2.text = "74.8"
    p_b2.font.name = FONT_FAMILY
    p_b2.font.size = Pt(32)
    p_b2.font.bold = True
    p_b2.font.color.rgb = C_GOLD
    p_b2.space_before = Pt(1)

    p_b3 = tf_cB.add_paragraph()
    p_b3.text = "80 Utility rating resolves the -35 deficit + Tier-1 IGL vocal calling. Directly lifts team win probability."
    p_b3.font.name = FONT_FAMILY
    p_b3.font.size = Pt(9.5)
    p_b3.font.color.rgb = C_TEXT_MUTED
    p_b3.space_before = Pt(2)

    add_footer(slide2, 2)

    set_notes(
        slide2,
        "Here is the core dilemma that every esports general manager faces today.",
        "Leaderboards rank players on kills and combat score, completely ignoring tactical role demands.",
        "In our PRD case study shown here, Voltz had the highest raw combat score (285 ACS), but scored a low 49.3 Roster Fit because 42% of his kills were useless exit frags during lost rounds.",
        "Mirage, despite a lower combat score, scored 74.8 because he solved Sentinels' exact tactical utility gap."
    )

    # =========================================================================
    # SLIDE 3 — STORYTELLING PIPELINE: SIMPLE CONNECTED FLOW & ENLARGED UI
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_bg(slide3)

    add_header(
        slide3,
        "A multi-layered intelligence pipeline from raw telemetry to roster fit.",
        "Evaluating combat, situational decision-making, role contribution, and cognitive biometrics."
    )

    # Top: Connected Component Flow (Crisp labels, generous spacing)
    flow_w = Inches(2.83)
    flow_h = Inches(1.50)
    flow_y = Inches(1.65)
    flow_gap = Inches(0.20)

    stages = [
        ("STAGE 01", "Match Ingestion", "Riot VCT & Steam CS2 APIs + 1,600 LAN observations.", C_TEXT_PRI),
        ("STAGE 02", "Role Calibration", "Normalizes by role; missing stats dropped, never zero.", C_CYAN),
        ("STAGE 03", "Cognitive Lab", "In-engine biometrics: 1560ms aim speed & spatial recall.", C_GOLD),
        ("STAGE 04", "Roster Decision", "PuLP ILP optimization under $1.5M payroll cap.", C_TEXT_PRI)
    ]
    for idx, (st_num, st_name, st_info, st_col) in enumerate(stages):
        s_left = Inches(0.70) + idx * (flow_w + flow_gap)
        add_card(slide3, s_left, flow_y, flow_w, flow_h)
        tb_st = slide3.shapes.add_textbox(s_left + Inches(0.12), flow_y + Inches(0.10), flow_w - Inches(0.24), flow_h - Inches(0.20))
        tf_st = tb_st.text_frame
        tf_st.word_wrap = True
        tf_st.margin_left = tf_st.margin_right = tf_st.margin_top = tf_st.margin_bottom = 0

        p1 = tf_st.paragraphs[0]
        p1.text = st_num
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = st_col

        p2 = tf_st.add_paragraph()
        p2.text = st_name
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_PRI
        p2.space_before = Pt(2)

        p3 = tf_st.add_paragraph()
        p3.text = st_info
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_MUTED
        p3.space_before = Pt(2)

    # Bottom: Enlarged Command Center & Gap Analyzer
    bot_y = Inches(3.35)
    bot_w = Inches(11.933)
    bot_h = Inches(3.40)
    add_card(slide3, Inches(0.70), bot_y, bot_w, bot_h)

    # Enlarged UI crop in left 62%
    img_w = Inches(7.30)
    img_h = bot_h - Inches(0.20)
    hub_crop = "SiteSnipts/02_intelligence_hub.png"
    if os.path.exists(hub_crop):
        slide3.shapes.add_picture(hub_crop, Inches(0.80), bot_y + Inches(0.10), img_w, img_h)

    # Short, high-impact guardrail labels in right 38%
    tb_guard = slide3.shapes.add_textbox(Inches(8.30), bot_y + Inches(0.25), Inches(4.10), bot_h - Inches(0.50))
    tf_g = tb_guard.text_frame
    tf_g.word_wrap = True
    tf_g.margin_left = tf_g.margin_right = tf_g.margin_top = tf_g.margin_bottom = 0

    p_gh = tf_g.paragraphs[0]
    p_gh.text = "Mathematical Guardrails"
    p_gh.font.name = FONT_FAMILY
    p_gh.font.size = Pt(14)
    p_gh.font.bold = True
    p_gh.font.color.rgb = C_GOLD

    guardrails = [
        ("Global Talent Index (GTI):", "Directional weighted sum; duelists weighted for entry frags, controllers for utility & survival."),
        ("Sample-Size Discounting:", "Candidates under 50 rounds are flagged Provisional. A short streak cannot outweigh 400 LAN rounds."),
        ("Double-Analyst Guardrail:", "Vocal IGL shotcalling scores remain locked until certified by two independent analyst annotations.")
    ]
    for lbl, desc in guardrails:
        p_item = tf_g.add_paragraph()
        p_item.text = lbl + " "
        p_item.font.name = FONT_FAMILY
        p_item.font.size = Pt(10)
        p_item.font.bold = True
        p_item.font.color.rgb = C_TEXT_PRI
        p_item.space_before = Pt(8)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    add_footer(slide3, 3)

    set_notes(
        slide3,
        "Now let's examine how Golden Snitch processes candidate data from raw match logs into roster decisions.",
        "We follow an objective 4-stage pipeline: ingestion, contextual normalization, cognitive biometrics, and integer programming.",
        "Our mathematical engine features strict sample-size discounting and double-analyst guardrails for in-game leaders.",
        "Next, let's examine the technical architecture and security guarantees that ensure enterprise trust."
    )

    # =========================================================================
    # SLIDE 4 — ARCHITECTURE & TRUST: SIMPLE CONNECTED FLOW & NO DENSE TEXT
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_bg(slide4)

    add_header(
        slide4,
        "Engineered for enterprise trust, auditability, and speed.",
        "Async FastAPI microservices, dual publisher telemetry, and immutable cryptographic audit trails."
    )

    # Left: 4-Layer Architecture Component Cards
    arch_left = Inches(0.70)
    arch_w = Inches(5.60)
    arch_top = Inches(1.65)
    layer_h = Inches(1.18)
    layer_gap = Inches(0.10)

    arch_stack = [
        ("LAYER 04 // CLIENT SURFACES", "React 18 + Vite, Tailwind CSS, Recharts 5-Axis Radars, WCAG AA", C_TEXT_PRI),
        ("LAYER 03 // SECURITY & GATEWAY", "JWT with PBKDF2-HMAC-SHA256, Server Token Revocation, Rate Limiting", C_CYAN),
        ("LAYER 02 // INTELLIGENCE & SCORING", "FastAPI Python 3.12, PuLP ILP Roster Optimizer, Percentile Engine", C_GOLD),
        ("LAYER 01 // DATA INGESTION & STORAGE", "HenrikDev Riot API, Valve Steam CS2 API, SQLite WAL Mode", C_TEXT_PRI)
    ]
    for idx, (l_title, l_desc, l_col) in enumerate(arch_stack):
        l_top = arch_top + idx * (layer_h + layer_gap)
        add_card(slide4, arch_left, l_top, arch_w, layer_h)
        tb_l = slide4.shapes.add_textbox(arch_left + Inches(0.15), l_top + Inches(0.12), arch_w - Inches(0.30), layer_h - Inches(0.24))
        tf_l = tb_l.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0

        p1 = tf_l.paragraphs[0]
        p1.text = l_title
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = l_col

        p2 = tf_l.add_paragraph()
        p2.text = l_desc
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(10)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_before = Pt(3)

    # Right: Data Governance Cards + Live DB Feed Screenshot
    gov_left = Inches(6.50)
    gov_w = Inches(6.133)
    gov_top = Inches(1.65)

    # Governance Highlights
    add_card(slide4, gov_left, gov_top, gov_w, Inches(2.20))
    tb_gov = slide4.shapes.add_textbox(gov_left + Inches(0.15), gov_top + Inches(0.12), gov_w - Inches(0.30), Inches(1.95))
    tf_g = tb_gov.text_frame
    tf_g.word_wrap = True
    tf_g.margin_left = tf_g.margin_right = tf_g.margin_top = tf_g.margin_bottom = 0

    p_gh = tf_g.paragraphs[0]
    p_gh.text = "Data Integrity & Governance Standards"
    p_gh.font.name = FONT_FAMILY
    p_gh.font.size = Pt(13)
    p_gh.font.bold = True
    p_gh.font.color.rgb = C_GOLD

    specs = [
        ("SHA-256 Decision Snapshots:", "Every recommendation run is sealed with a cryptographic hash for board auditing."),
        ("CSV Formula Injection Defense:", "Sanitizes uploads by stripping =, @, +, - triggers to block spreadsheet attacks."),
        ("Graceful Dual-Telemetry Fallback:", "Auto-falls back to verified offline LAN datasets if live publisher APIs throttle.")
    ]
    for lbl, desc in specs:
        p_item = tf_g.add_paragraph()
        p_item.text = lbl + " "
        p_item.font.name = FONT_FAMILY
        p_item.font.size = Pt(9.5)
        p_item.font.bold = True
        p_item.font.color.rgb = C_TEXT_PRI
        p_item.space_before = Pt(3)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    # Bottom Right: Enlarged Live DB Feed Screenshot
    add_card(slide4, gov_left, gov_top + Inches(2.35), gov_w, Inches(2.65))
    db_img = "SiteSnipts/03_player_database.png"
    if os.path.exists(db_img):
        slide4.shapes.add_picture(db_img, gov_left + Inches(0.08), gov_top + Inches(2.43), gov_w - Inches(0.16), Inches(2.49))

    add_footer(slide4, 4)

    set_notes(
        slide4,
        "Enterprise esports franchises require high reliability, data protection, and auditable governance.",
        "Golden Snitch is built on an asynchronous FastAPI backend and React 18 client with SQLite WAL mode.",
        "We implement cryptographic SHA-256 decision snapshots, formula injection sanitization, and graceful offline fallback.",
        "Now let's examine how this translates into player development and roster optimization."
    )

    # =========================================================================
    # SLIDE 5 — PLAYER DEVELOPMENT & TEAM ROI: RESTRAINED 50/50 COMPOSITION
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_bg(slide5)

    add_header(
        slide5,
        "Aligning individual player potential with franchise roster needs.",
        "Counterfactual synergy simulation, cognitive hardware profiling, and multi-roster management."
    )

    split_w = Inches(5.80)
    split_h = Inches(5.05)
    split_top = Inches(1.65)

    # Left: Player Hardware & Cognitive Profiling
    add_card(slide5, Inches(0.70), split_top, split_w, split_h)
    tb_pl = slide5.shapes.add_textbox(Inches(0.85), split_top + Inches(0.12), split_w - Inches(0.30), Inches(1.65))
    tf_pl = tb_pl.text_frame
    tf_pl.word_wrap = True
    tf_pl.margin_left = tf_pl.margin_right = tf_pl.margin_top = tf_pl.margin_bottom = 0

    p_lh = tf_pl.paragraphs[0]
    p_lh.text = "Individual Biometrics & Reflex Hardware"
    p_lh.font.name = FONT_FAMILY
    p_lh.font.size = Pt(13)
    p_lh.font.bold = True
    p_lh.font.color.rgb = C_CYAN

    p_lb = tf_pl.add_paragraph()
    p_lb.text = "In-engine standardized biometrics benchmark player hardware under pressure:"
    p_lb.font.name = FONT_FAMILY
    p_lb.font.size = Pt(10)
    p_lb.font.color.rgb = C_TEXT_MUTED
    p_lb.space_before = Pt(2)

    biometrics = [
        "Aim Acquisition: 1560.43 ms targeting speed with 69.77% hit accuracy across 30 targets.",
        "Reaction Latency: Visual stimulus click response measured across 5 scored trials.",
        "Spatial Working Memory: Progressive 3x3 tactical grid pattern retention."
    ]
    for pt in biometrics:
        p_pt = tf_pl.add_paragraph()
        p_pt.text = "• " + pt
        p_pt.font.name = FONT_FAMILY
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = C_TEXT_SEC
        p_pt.space_before = Pt(2)

    # Enlarged Cognitive Lab Screenshot
    cog_img = "SiteSnipts/04_cognitive_lab.png"
    if os.path.exists(cog_img):
        slide5.shapes.add_picture(cog_img, Inches(0.85), split_top + Inches(1.88), split_w - Inches(0.30), Inches(3.02))

    # Right: Franchise ROI & Operational Metrics
    right_x = Inches(6.83)
    add_card(slide5, right_x, split_top, split_w, split_h)

    tb_pr = slide5.shapes.add_textbox(right_x + Inches(0.15), split_top + Inches(0.12), split_w - Inches(0.30), split_h - Inches(0.24))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True
    tf_pr.margin_left = tf_pr.margin_right = tf_pr.margin_top = tf_pr.margin_bottom = 0

    p_rh = tf_pr.paragraphs[0]
    p_rh.text = "Franchise Roster Optimization & Value"
    p_rh.font.name = FONT_FAMILY
    p_rh.font.size = Pt(13)
    p_rh.font.bold = True
    p_rh.font.color.rgb = C_GOLD

    p_rb = tf_pr.add_paragraph()
    p_rb.text = "Demonstrated operational improvements across competitive tournament rosters:"
    p_rb.font.name = FONT_FAMILY
    p_rb.font.size = Pt(10)
    p_rb.font.color.rgb = C_TEXT_MUTED
    p_rb.space_before = Pt(2)

    kpis = [
        ("88% TIME SAVINGS", "Scouting cycle reduced from 18.5 hours of manual sheets to 2.2 hours via automated gap matching."),
        ("+12.8% ECO CONVERSION", "Teams signing candidates who fill diagnosed utility deficits showed +12.8% higher conversion on low-buy rounds."),
        ("ZERO DECOY RISK", "Mathematical filtering eliminates expensive buyout errors on pad-stat fraggers whose frags carry no win equity."),
        ("GRASSROOTS EQUITY", "Democratizes talent discovery for Tier-2/3 athletes in Maharashtra and India whose support value is hidden on public sites.")
    ]
    for idx, (k_title, k_desc) in enumerate(kpis):
        p_kt = tf_pr.add_paragraph()
        p_kt.text = k_title
        p_kt.font.name = FONT_FAMILY
        p_kt.font.size = Pt(11)
        p_kt.font.bold = True
        p_kt.font.color.rgb = C_TEXT_PRI if idx != 1 else C_GOLD
        p_kt.space_before = Pt(9)

        p_kd = tf_pr.add_paragraph()
        p_kd.text = k_desc
        p_kd.font.name = FONT_FAMILY
        p_kd.font.size = Pt(9.5)
        p_kd.font.color.rgb = C_TEXT_MUTED
        p_kd.space_before = Pt(1)

    add_footer(slide5, 5)

    set_notes(
        slide5,
        "Beyond statistical scoring, Golden Snitch bridges player hardware biometrics with franchise roster optimization.",
        "Our Cognitive Lab evaluates human reaction latency, aim precision, and spatial working memory under competitive conditions.",
        "On the franchise side, our counterfactual simulations deliver an 88% reduction in scouting cycle time and eliminate expensive buyout mistakes.",
        "Finally, let us review our product roadmap and strategic vision."
    )

    # =========================================================================
    # SLIDE 6 — ROADMAP & VALUE PROPOSITION: CLEAN TIMELINE & TAGLINE
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_bg(slide6)

    add_header(
        slide6,
        "From production-ready MVP to the industry standard in esports ops.",
        "Verified on 1,726 tournament matches, with a clear trajectory for computer-vision and mobile expansion."
    )

    # 4 Milestone Flow Panels
    m_w = Inches(2.83)
    m_h = Inches(3.20)
    m_top = Inches(1.65)
    m_gap = Inches(0.20)

    roadmap = [
        ("CURRENT STATE", "Operational Solution",
         "• Full scoring & gap matching engine\n• 1,726 VCT matches ingested\n• PuLP ILP salary cap optimizer\n• 112 passing pytest unit tests\n• Deployed live on Railway cloud", C_GOLD),
        ("Q1–Q2 2027", "Regional Pilot Cohort",
         "• 60-day closed beta with 3-5 collegiate & semi-pro academies\n• Metric calibration on regional Indian LAN tournaments\n• Coach feedback iteration", C_TEXT_PRI),
        ("Q3 2027", "Broadcast CV Parser",
         "• Automated minimap computer vision for match broadcast VODs\n• Automated flash angle & smoke micro-telemetry extraction\n• Eliminates manual sheet uploads", C_CYAN),
        ("Q4 2027", "Contract API & Mobile PWA",
         "• Publisher contract database hooks with buyout alert feeds\n• Mobile PWA for live venue floor scouting by coaches\n• Multi-title expansion: BGMI / Dota", C_TEXT_PRI)
    ]
    for idx, (m_period, m_title, m_desc, m_col) in enumerate(roadmap):
        m_left = Inches(0.70) + idx * (m_w + m_gap)
        add_card(slide6, m_left, m_top, m_w, m_h)
        tb_m = slide6.shapes.add_textbox(m_left + Inches(0.15), m_top + Inches(0.12), m_w - Inches(0.30), m_h - Inches(0.24))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True
        tf_m.margin_left = tf_m.margin_right = tf_m.margin_top = tf_m.margin_bottom = 0

        p1 = tf_m.paragraphs[0]
        p1.text = m_period
        p1.font.name = FONT_FAMILY
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = m_col

        p2 = tf_m.add_paragraph()
        p2.text = m_title
        p2.font.name = FONT_FAMILY
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_TEXT_PRI
        p2.space_before = Pt(3)

        p3 = tf_m.add_paragraph()
        p3.text = m_desc
        p3.font.name = FONT_FAMILY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_TEXT_MUTED
        p3.space_before = Pt(6)

    # Bottom Editorial Statement Card
    banner_y = Inches(5.08)
    banner_w = Inches(11.933)
    banner_h = Inches(1.55)
    add_card(slide6, Inches(0.70), banner_y, banner_w, banner_h, bg=C_SURFACE_LIGHT)

    tb_tag = slide6.shapes.add_textbox(Inches(0.90), banner_y + Inches(0.20), banner_w - Inches(0.40), banner_h - Inches(0.40))
    tf_tag = tb_tag.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0

    p_th = tf_tag.paragraphs[0]
    p_th.text = "“Scout beyond the scoreboard.”"
    p_th.font.name = FONT_FAMILY
    p_th.font.size = Pt(24)
    p_th.font.bold = True
    p_th.font.color.rgb = C_GOLD

    p_td = tf_tag.add_paragraph()
    p_td.text = "Golden Snitch replaces subjective guesswork with objective, auditable decision intelligence built around real roster needs."
    p_td.font.name = FONT_FAMILY
    p_td.font.size = Pt(12)
    p_td.font.color.rgb = C_TEXT_PRI
    p_td.space_before = Pt(4)

    add_footer(slide6, 6)

    set_notes(
        slide6,
        "In conclusion, Golden Snitch is not a theoretical concept; it is an operational solution.",
        "We have verified our mathematical scoring against 1,726 VCT matches, validated 112 automated tests, and proven our ability to detect statistical decoys.",
        "Our roadmap takes us through collegiate academy pilots, computer-vision broadcast ingestion, and multi-game expansion.",
        "We invite you to test our live deployment and join us in scouting beyond the scoreboard. Thank you."
    )

    prs.save(target_path)
    print(f"Presentation generated successfully: {target_path}")

if __name__ == "__main__":
    build_presentation()
