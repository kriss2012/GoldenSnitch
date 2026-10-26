import pptx

prs = pptx.Presentation('Golden Snitch (ScoutIQ) – Pitch Deck.pptx')
print(f"Presentation has {len(prs.slides)} slides. Dimensions: {prs.slide_width} x {prs.slide_height}")
for idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {idx+1} (Layout: {slide.slide_layout.name}) ====================")
    for s_idx, shape in enumerate(slide.shapes):
        print(f"Shape {s_idx}: name='{shape.name}', type={shape.shape_type}, pos=({shape.left}, {shape.top}), size=({shape.width}, {shape.height})")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                txt = p.text.strip().encode('ascii', errors='replace').decode('ascii')
                if txt:
                    print(f"   P: '{txt}'")
