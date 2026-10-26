import pptx

prs = pptx.Presentation('Golden Snitch (ScoutIQ) – Pitch Deck.pptx')
slide = prs.slides[0]
print(f"Slide width: {prs.slide_width.inches} inches, height: {prs.slide_height.inches} inches")
for s in slide.shapes:
    print(f"Shape: {s.name}")
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                font = r.font
                color = font.color.rgb if font.color and font.color.type == pptx.enum.dml.MSO_COLOR_TYPE.RGB else font.color.theme_color if font.color else None
                print(f"   Run: '{r.text}' | Font: {font.name}, Size: {font.size}, Color: {color}, Bold: {font.bold}")
