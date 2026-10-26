import pptx

prs = pptx.Presentation("Golden Snitch (ScoutIQ) – Pitch Deck.pptx")
print(f"Total slides: {len(prs.slides)}")
print(f"Slide dimensions: {prs.slide_width.inches} x {prs.slide_height.inches} inches")

for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    shapes = list(slide.shapes)
    print(f"Number of shapes: {len(shapes)}")
    images = [s for s in shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE]
    text_boxes = [s for s in shapes if s.has_text_frame]
    print(f"Images: {len(images)} | Text boxes: {len(text_boxes)}")
    for s in shapes:
        # Check boundary
        left_in = s.left / 914400.0
        top_in = s.top / 914400.0
        width_in = s.width / 914400.0
        height_in = s.height / 914400.0
        right_in = left_in + width_in
        bottom_in = top_in + height_in
        if right_in > 13.35 or bottom_in > 7.55:
            print(f"  WARNING: Shape {s.name} extends out of bounds! Right={right_in:.2f}, Bottom={bottom_in:.2f}")
    # Print slide title text
    if text_boxes:
        first_text = [p.text for p in text_boxes[0].text_frame.paragraphs if p.text.strip()]
        second_text = [p.text for p in text_boxes[1].text_frame.paragraphs if p.text.strip()] if len(text_boxes) > 1 else []
        print(f"  Header: {first_text} | {second_text}")
