import os
import pptx

files = [
    'Golden Snitch (ScoutIQ) – Pitch Deck.pptx',
    'Golden_Snitch_ScoutIQ.pptx',
    'Golden_Snitch_ScoutIQ (1).pptx'
]

for name in files:
    if os.path.exists(name):
        prs = pptx.Presentation(name)
        print(f"\n============================\n{name}: {len(prs.slides)} slides\n============================")
        for i, s in enumerate(prs.slides):
            texts = [p.text.strip() for shape in s.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs if p.text.strip()]
            img_count = sum(1 for shape in s.shapes if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE)
            print(f"Slide {i+1}:")
            print(f"  Titles/Texts ({len(texts)} items): {texts[:4]}")
            print(f"  Pictures: {img_count}")
