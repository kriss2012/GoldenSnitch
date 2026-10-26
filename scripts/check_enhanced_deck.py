import pptx

prs = pptx.Presentation("Golden_Snitch_Pitch_Deck_Enhanced.pptx")
print(f"Presentation has {len(prs.slides)} slides. Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")

for i, slide in enumerate(prs.slides):
    shapes = list(slide.shapes)
    images = [s for s in shapes if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE]
    text_boxes = [s for s in shapes if s.has_text_frame]
    print(f"\n==================== SLIDE {i+1} ====================")
    print(f"Total Shapes: {len(shapes)} | Images: {len(images)} | Text Boxes: {len(text_boxes)}")
    for s in shapes:
        left = s.left / 914400.0
        top = s.top / 914400.0
        w = s.width / 914400.0
        h = s.height / 914400.0
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE]: {s.name} at pos=({left:.2f}, {top:.2f}), size=({w:.2f}, {h:.2f})")
        elif s.has_text_frame:
            first_p = s.text_frame.paragraphs[0].text.strip()
            clean_p = first_p[:60].encode('ascii', errors='replace').decode('ascii')
            if clean_p:
                print(f"  [TEXT]: '{clean_p}' at pos=({left:.2f}, {top:.2f})")
